"""Controlled, offline import of reviewed government workbooks (Python/openpyxl).

Usage: python scripts/import-official-statistics.py /absolute/source-directory
Expected files: abs.xlsx, sales.xlsx, crime.xlsx, seifa.xlsx, approvals.zip,
schools.csv. Original source URLs and hashes are recorded in import-manifest.json.
Downloads are never committed.
No database writes; inspect the JSON diff, run tests, then release normally.
"""
import hashlib
import json
import pathlib
import sys
import csv
import io
import zipfile

import openpyxl

root = pathlib.Path(__file__).resolve().parents[1]
dest = root / 'src/lib/layers/data'
source = pathlib.Path(sys.argv[1])
crosswalk = json.loads((dest / 'vic-lga-crosswalk.json').read_text())['rows']
vgv = {x['vgv_name']: x['code'] for x in crosswalk}
csa = {x['csa_name']: x['code'] for x in crosswalk}


def number(value):
    return value if isinstance(value, (float, int)) and not isinstance(value, bool) else None


def save(name, data):
    (dest / (name + '.json')).write_text(json.dumps(data, separators=(',', ':'), allow_nan=False) + '\n')


def workbook(name):
    return openpyxl.load_workbook(source / (name + '.xlsx'), read_only=True, data_only=True)


population = {}
for row in workbook('abs')['Table 1'].iter_rows(min_row=7, max_col=35, values_only=True):
    if not isinstance(row[8], int):
        continue
    code = str(row[8])
    assert code not in population
    population[code] = {'name': row[9], 'erp': [number(v) for v in row[10:35]]}
assert len(population) > 2400
save('population', population)

rows = list(workbook('sales')['Table 1'].iter_rows(max_col=54, values_only=True))
sales = {}
for i, row in enumerate(rows):
    if not str(row[0] or '').startswith('Residential price statistics'):
        continue
    name = rows[i - 1][0].strip()
    if name in ['Melbourne Metropolitan Area', 'Country Victoria', 'Victoria', 'French Island']:
        continue  # Explicitly excluded aggregates / non-LGA island, never assigned to a municipality.
    code = vgv[name]  # Fail on unknown identities; no fuzzy/name joins at runtime.
    headers = rows[i + 2]
    years = [j for j, value in enumerate(headers) if value == 'Year']
    medians = [j for j, value in enumerate(headers) if value == 'Median Price']
    counts = [j for j, value in enumerate(headers) if value == 'No. of Sales']
    assert len(years) == 1 and len(medians) == len(counts) == 3
    values = {}
    for r in rows[i + 3:i + 35]:
        year = r[years[0]]
        if not isinstance(year, int) or not 2015 <= year <= 2025:
            continue  # Exclude the workbook's partial 2026 observations.
        record = {}
        for kind, price, count in zip(['house', 'unit', 'land'], medians, counts):
            n, value = number(r[count]), number(r[price])
            record[kind] = value if n and value and value > 0 else None
            record[kind + '_sales'] = n
        values[str(year)] = record
    assert len(values) == 11
    sales[code] = values
assert len(sales) == 79
save('vic-sales', sales)

crime = {}
for row in workbook('crime')['Table 01'].iter_rows(min_row=2, max_col=6, values_only=True):
    year, ending, _, name, count, rate = row
    if not year or not name:
        continue
    name = name.strip()
    if name in ['Justice Institutions and Immigration Facilities', 'Unincorporated Vic', 'Total']:
        continue
    code = csa[name]
    assert ending == 'June'
    crime.setdefault(code, {})[str(year)] = {'offences': number(count), 'rate': number(rate)}

for row in workbook('crime')['Table 02'].iter_rows(min_row=2, max_col=10, values_only=True):
    year, ending, _, name, _, subdivision, subgroup, count, _, _ = row
    if not year or not name or name.strip() not in csa:
        continue
    code = csa[name.strip()]
    metrics = []
    if str(subdivision).startswith('A20 '):
        metrics.append('assault')
    if str(subdivision).startswith('B30 '):
        metrics.append('burglary')
    if str(subdivision).startswith('B40 '):
        metrics.append('theft')
    if str(subgroup).startswith('B41 '):
        metrics.append('motor_theft')
    for metric in metrics:
        target = crime[code][str(year)]
        target[metric] = target.get(metric, 0) + count
assert len(crime) == 79 and all(len(x) == 10 for x in crime.values())
save('vic-crime', crime)

seifa = {}
for row in workbook('seifa')['Table 1'].iter_rows(min_row=7, max_col=11, values_only=True):
    if not isinstance(row[0], int):
        continue
    code = str(row[0])
    seifa[code] = {key: number(row[col]) for key, col in [
        ('irsd', 2), ('irsd_decile', 3), ('irsad', 4), ('irsad_decile', 5),
        ('ier', 6), ('ier_decile', 7), ('ieo', 8), ('ieo_decile', 9)]}
assert len(seifa) > 2300
save('seifa', seifa)

# Use source total-sector rows once, never add totals to their components.
# Complete FY only; suppress the annual sum if any monthly value is missing.
approvals = {}
with zipfile.ZipFile(source / 'approvals.zip') as archive:
    for row in csv.DictReader(io.StringIO(archive.read('sa2_2526_formatted.csv').decode('utf-8-sig'))):
        code = row['sa2_code']
        if len(code) != 9 or code not in population or row['type_work'] != '1' or row['own_sector'] != '9':
            continue
        kind = {'110': 'houses', '150': 'other_residential', '100': 'new_dwellings'}.get(row['type_bld'])
        if not kind:
            continue
        approvals.setdefault(code, {}).setdefault(kind, []).append(float(row['dwl']) if row['dwl'].replace('.', '', 1).isdigit() else None)
for code, values in approvals.items():
    for kind, months in list(values.items()):
        values[kind] = sum(months) if len(months) == 12 and all(v is not None for v in months) else None
assert len(approvals) > 2400
save('building-approvals', approvals)

schools = []
with (source / 'schools.csv').open(encoding='cp1252', newline='') as file:
    for row in csv.DictReader(file):
        x, y = float(row['X']), float(row['Y'])
        assert 140 < x < 151 and -40 < y < -33
        fields = {k: row[k] for k in ['School_No', 'School_Name', 'School_Type', 'School_Status', 'Education_Sector', 'Address_Line_1', 'Address_Line_2', 'Address_Town', 'Address_Postcode', 'LGA_Name']}
        fields['location_id'] = ':'.join(row[k] for k in ['School_No', 'Education_Sector', 'Entity_Type', 'X', 'Y'])
        schools.append({'type': 'Feature', 'properties': fields, 'geometry': {'type': 'Point', 'coordinates': [x, y]}})
assert len(schools) > 2000 and len({f['properties']['location_id'] for f in schools}) == len(schools)
save('vic-schools', schools)

urls = {
    'abs': 'https://www.abs.gov.au/statistics/people/population/regional-population/2024-25/32180DS0003_2001-25.xlsx',
    'sales': 'https://www.land.vic.gov.au/__data/assets/excel_doc/0032/775616/year-summary-2025.xlsx',
    'crime': 'https://files.crimestatistics.vic.gov.au/2026-09/Data_Tables_LGA_Recorded_Offences_Year_Ending_June_2026.xlsx',
    'seifa': 'https://www.abs.gov.au/statistics/people/people-and-communities/socio-economic-indexes-areas-seifa-australia/2021/Statistical%20Area%20Level%202%2C%20Indexes%2C%20SEIFA%202021.xlsx',
    'approvals': 'https://www.abs.gov.au/statistics/industry/building-and-construction/building-approvals-australia/aug-2026/Statistical%20Area%202_Australia_2025-26.zip',
    'schools': 'https://content.vic.gov.au/DE-School-Locations%20-2026',
}
save('import-manifest', {'verified_on': '2026-10-31', 'sources': [
    {'id': k, 'url': u, 'sha256': hashlib.sha256((source / (k + ('.zip' if k == 'approvals' else '.csv' if k == 'schools' else '.xlsx'))).read_bytes()).hexdigest(),
     'licence': 'CC BY 4.0', 'licence_evidence': {
         'abs': 'https://www.abs.gov.au/website-privacy-copyright-and-disclaimer',
         'sales': 'https://discover.data.vic.gov.au/dataset/victorian-property-sales-report-time-series',
         'crime': 'https://discover.data.vic.gov.au/dataset/data-tables-recorded-offences',
         'seifa': 'https://www.abs.gov.au/website-privacy-copyright-and-disclaimer',
         'approvals': 'https://www.abs.gov.au/website-privacy-copyright-and-disclaimer',
         'schools': 'https://discover.data.vic.gov.au/dataset/school-locations-2026',
     }[k]} for k, u in urls.items()]})
print(f'Imported {len(population)} SA2s, {len(sales)} sales LGAs and {len(crime)} crime LGAs')
