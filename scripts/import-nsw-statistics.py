"""Controlled offline NSW aggregate import; never reads individual sales records.

Usage: python scripts/import-nsw-statistics.py /absolute/source-directory
Inputs: nsw-sales.xlsx, nsw-rent.xlsx, nsw-crime.xlsm,
nsw-lga-identities.json (ABS 2025), nsw-lga-identities-2016.json (ABS 2016).
Read XLSM cells only; never execute macros. No database or network writes.
Review original notes/licences, generated crosswalk and JSON diff before release.
"""
import datetime
import hashlib
import json
import pathlib
import re
import sys

import openpyxl

source = pathlib.Path(sys.argv[1])
dest = pathlib.Path(__file__).resolve().parents[1] / 'src/lib/layers/data'

# Period labels in market notes are embedded images, reviewed manually. Lock
# reviewed originals so a newer workbook cannot silently inherit these dates.
expected_sha256 = {'nsw-sales.xlsx': '6a72f007c898e0dbf9ceabf654f44ef0df80d8280d845818c88edbdd83431e18', 'nsw-rent.xlsx': '34bfd9add8fdd1e41ee6c91eb8cb2714854ec6a4da219a0d99e8a43c44c35bec', 'nsw-crime.xlsm': '9b13b9e878875b01a80d27703dc858c385ae693b5c929539079e135202e9484f', 'nsw-lga-identities.json': '763e400943d0e4fdc07fecb5a00e896eb98ad35dec8b0520c166bc1b343ec01b', 'nsw-lga-identities-2016.json': '3b6175532d696f090bbcda5f3934ed133681f19606ad12fd47c4c2f5422bf399'}
for name, digest in expected_sha256.items():
    assert hashlib.sha256((source / name).read_bytes()).hexdigest() == digest, f'Review the new original and period before importing: {name}'



def save(name, value):
    (dest / (name + '.json')).write_text(json.dumps(value, separators=(',', ':'), allow_nan=False) + '\n')


def numeric(value):
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return value
    text = str(value).strip().replace(',', '')
    if not re.fullmatch(r'-?(?:\d+(?:\.\d*)?|\.\d+)%?', text):
        return None
    return float(text.rstrip('%')) if '.' in text or text.endswith('%') else int(text)


def observation(value, scale=1):
    number = numeric(value)
    return {'value': number * scale if number is not None else None,
            'source_marker': None if number is not None or value is None else str(value).strip()}


def workbook(name):
    return openpyxl.load_workbook(source / name, read_only=True, data_only=True, keep_vba=False)


old_rows = [f['attributes'] for f in json.loads((source / 'nsw-lga-identities-2016.json').read_text())['features']]
new_rows = [f['attributes'] for f in json.loads((source / 'nsw-lga-identities.json').read_text())['features']]
old_by_name = {r['lga_name_2016'].casefold(): r for r in old_rows}
new_by_code = {r['LGA_CODE_2025']: r for r in new_rows}
new_by_name = {r['LGA_NAME_2025'].casefold(): r for r in new_rows}
assert len(old_by_name) == len(old_rows) == 130
assert len(new_by_code) == len(new_by_name) == len(new_rows) == 131

# Explicit NSW-state qualifications, not a generic removal of disambiguators.
qualified = {'Bayside': '10500', 'Campbelltown': '11500', 'Central Coast': '11650'}
# ABS ASGS 2018 Appendix 1 and NSW Electoral Commission record these renames.
# These are identity changes, not a population-weighted conversion of medians.
renamed_codes = {'13510': '12160', '18230': '12390'}
rename_evidence = [
    'https://www.abs.gov.au/ausstats/abs%40.nsf/Previousproducts/3218.0Appendix12017-18?issue=2017-18&num=&opendocument=&prodno=3218.0&tabname=Notes&view=',
    'https://elections.nsw.gov.au/elections/find-my-electorate/councils/cootamundra-gundagai',
    'https://elections.nsw.gov.au/elections/find-my-electorate/councils/dubbo',
]


def identity(name):
    if name in qualified:
        code = qualified[name]
        return {'source_name': name, 'code': code, 'name': new_by_code[code]['LGA_NAME_2025'],
                'mapping': 'Explicit NSW-qualified ABS 2025 identity'}
    if name.casefold() in new_by_name:
        row = new_by_name[name.casefold()]
        return {'source_name': name, 'code': row['LGA_CODE_2025'], 'name': row['LGA_NAME_2025'],
                'mapping': 'Exact unique publisher label to ABS 2025 identity'}
    # The original DCJ labels match ABS 2016 names without their legal type.
    # Enumerate the types; retain original full names and codes in this crosswalk.
    candidates = [old_by_name[label.casefold()] for suffix in (' (A)', ' (C)', ' (M)', ' (R)', ' (S)')
                  if (label := name + suffix).casefold() in old_by_name]
    assert len(candidates) <= 1, name
    if candidates:
        old = candidates[0]
        old_code = old['lga_code_2016']
        code = renamed_codes.get(old_code, old_code)
        if code in new_by_code:
            return {'source_name': name, 'source_abs_name_2016': old['lga_name_2016'],
                    'source_abs_code_2016': old_code, 'code': code,
                    'name': new_by_code[code]['LGA_NAME_2025'],
                    'mapping': 'Reviewed ABS rename' if code != old_code else 'Same official LGA code across 2016 and 2025 identities'}
        return {'source_name': name, 'source_abs_name_2016': old['lga_name_2016'],
                'source_abs_code_2016': old_code, 'code': None,
                'reason': 'Retired source LGA code; boundary correspondence not verified. Never assigned by similar name.'}
    return {'source_name': name, 'code': None, 'reason': 'No verified bounded ABS LGA identity; not assigned to another region.'}


sales = {}
sales_keys = set()
dcj_names = set()
for row in workbook('nsw-sales.xlsx')['LGA'].iter_rows(min_row=8, values_only=True):
    if row[:3] != ('Total', 'Total', 'Total') or row[3] in (None, 'Total'):
        continue  # The workbook repeats LGAs under several regional rollups.
    name, kind = row[3:5]
    assert kind in ('Total', 'Non Strata', 'Strata')
    assert (name, kind) not in sales_keys
    sales_keys.add((name, kind)); dcj_names.add(name)
    mapped = identity(name)
    if not mapped['code']:
        continue
    code = mapped['code']
    record = sales.setdefault(code, {'name': mapped['name'], 'source_name': name, '2026Q1': {}})
    record['2026Q1'][kind] = {
        'median': observation(row[6], 1000), 'sales': observation(row[9]),
        'quarterly_change': observation(row[10]), 'annual_change': observation(row[11]),
    }
assert len(dcj_names) == 127 and len(sales_keys) == 341
assert {name for name in dcj_names if not identity(name)['code']} == {'Armidale Regional'}
assert len(sales) == 126
save('nsw-sales', sales)

rent = {}
rent_keys = set()
rent_names = set()
for row in workbook('nsw-rent.xlsx')['LGA'].iter_rows(min_row=10, values_only=True):
    if row[:3] != ('Total', 'Total', 'Total') or row[3] in (None, 'Total'):
        continue
    name, kind, bedrooms = row[3:6]
    key = (name, kind, bedrooms)
    assert key not in rent_keys
    rent_keys.add(key); rent_names.add(name)
    mapped = identity(name)
    if not mapped['code']:
        continue
    code = mapped['code']
    record = rent.setdefault(code, {'name': mapped['name'], 'source_name': name, '2026Q2': {}})
    record['2026Q2'][kind + ' / ' + bedrooms] = {
        'median': observation(row[7]), 'new_bonds': observation(row[9]),
        'total_bonds_held': observation(row[10]), 'quarterly_change': observation(row[11]),
        'annual_change': observation(row[12]),
    }
assert rent_names - dcj_names == {'Balranald'} and not (dcj_names - rent_names)
assert len(rent_names) == 128 and len(rent_keys) == 3937 and len(rent) == 127
save('nsw-rent', rent)

crime = {}
crime_names = set()
crime_keys = set()
sheet = workbook('nsw-crime.xlsm')['Data']
rows = sheet.values
header = next(rows)
dates = list(header[3:])
assert dates[0] == datetime.datetime(1995, 1, 1) and dates[-1] == datetime.datetime(2026, 6, 1)
categories = {
    'Break and enter dwelling': 'break_enter_dwelling',
    'Break and enter non-dwelling': 'break_enter_non_dwelling',
    'Motor vehicle theft': 'motor_vehicle_theft',
    'Domestic violence related assault': 'domestic_violence_assault',
    'Non-domestic violence related assault': 'non_domestic_violence_assault',
}
for row in rows:
    name, group, kind = row[:3]
    assert (name, group, kind) not in crime_keys
    crime_keys.add((name, group, kind)); crime_names.add(name)
    mapped = identity(name)
    if not mapped['code'] or kind not in categories:
        continue
    record = crime.setdefault(mapped['code'], {'name': mapped['name'], 'source_name': name})
    for year in range(2017, 2027):
        indices = [i for i, date in enumerate(dates, 3)
                   if datetime.datetime(year - 1, 7, 1) <= date < datetime.datetime(year, 7, 1)]
        assert len(indices) == 12
        values = [numeric(row[i]) for i in indices]
        value = sum(values) if all(v is not None for v in values) else None
        if value is not None:
            assert value >= 0 and int(value) == value
        record.setdefault(str(year), {})[categories[kind]] = value
assert len(crime_names) == 131 and len(crime_keys) == 131 * 62
assert len(crime) == 128
assert {n for n in crime_names if not identity(n)['code']} == {'In Custody', 'Lord Howe Island', 'Unincorporated Far West'}
save('nsw-crime', crime)

save('nsw-lga-crosswalk', {
    'version': '31 October 2026; ABS LGA 2016/2025 identities',
    'boundary_use': 'ABS 2025 reference polygons; source workbooks label LGAs but do not specify ASGS boundary year. No historical boundary reconstruction or areal redistribution.',
    'rename_evidence': rename_evidence,
    'dcj': [{**identity(n), 'sales_source_row': n in dcj_names, 'rent_source_row': n in rent_names} for n in sorted(dcj_names | rent_names)],
    'bocsar': [identity(n) for n in sorted(crime_names)],
})

sources = {
    'nsw-sales.xlsx': 'https://dcj.nsw.gov.au/documents/about-us/families-and-communities-statistics/housing-and-rent-sales/sales-tables-march-2026-quarter.xlsx',
    'nsw-rent.xlsx': 'https://dcj.nsw.gov.au/documents/about-us/families-and-communities-statistics/housing-and-rent-sales/rent-tables-june-2026-quarter.xlsx',
    'nsw-crime.xlsm': 'https://bocsarblob.blob.core.windows.net/bocsar-open-data/RCI_offencebymonth.xlsm',
    'nsw-lga-identities.json': "https://geo.abs.gov.au/arcgis/rest/services/ASGS2025/LGA/MapServer/1/query?f=json&where=STATE_CODE_2021%3D%271%27&outFields=LGA_CODE_2025%2CLGA_NAME_2025&returnGeometry=false&resultRecordCount=1000",
    'nsw-lga-identities-2016.json': "https://geo.abs.gov.au/arcgis/rest/services/ASGS2016/LGA/MapServer/1/query?f=json&where=state_code_2016%3D%271%27&outFields=lga_code_2016%2Clga_name_2016&returnGeometry=false&resultRecordCount=1000",
}
save('nsw-import-manifest', {
    'version': 'NSW DCJ Report 156; BOCSAR June 2026; controlled import 31 October 2026',
    'sources': [{'file': name, 'url': url, 'bytes': (source / name).stat().st_size,
                 'sha256': hashlib.sha256((source / name).read_bytes()).hexdigest()} for name, url in sources.items()],
    'dcj_licence': 'Creative Commons Attribution 4.0; original notes Crown copyright. No separate noncommercial exception in the downloaded aggregate workbooks.',
    'dcj_licence_evidence': 'https://dcj.nsw.gov.au/statements/copyright-and-disclaimer.html',
    'bocsar_licence': 'Creative Commons Attribution, as explicitly assigned to the exact external-hosted resource by Data.NSW; original portal licence-family link retained.',
    'bocsar_licence_url': 'https://opendefinition.org/licenses/cc-by/',
    'bocsar_licence_evidence': 'https://data.nsw.gov.au/data/dataset/nsw-crime-and-policing-statistics/resource/1d5b2851-52e9-4327-a81b-19149c63f736',
    'boundary_licence_evidence': 'https://www.abs.gov.au/website-privacy-copyright-and-disclaimer',
    'reuse': 'Commercial reuse and adapted aggregates under applicable CC BY attribution; preserve source dates, indicate modifications, no endorsement. Public geometry only cached; private workspace data never included.',
    'geography': '126/127 DCJ sales LGAs and 127/128 rent LGAs mapped; Armidale Regional retained unmatched. Balranald has no original sales row and stays missing. All 128 bounded BOCSAR council labels mapped; three non-council/special source labels excluded, never assigned to neighbours.',
    'periods': {'sales': 'March quarter 2026; extracted 31 July 2026; contract date, not settlement date',
                'rent': 'June quarter 2026; extracted 31 July 2026; new rental bonds, not all existing tenancies or asking rents',
                'crime': 'Ten rolling 12-month periods ending June 2017–2026, sums of original monthly incident counts'},
    'refresh': 'Review new quarterly originals, notes, licence and identities; rerun this offline importer, diff hashes/values/unmatched rows, verify and release. No unattended ingestion.',
    'limitations': ['s/- suppression preserved, never zero. Samples of 11–30 carry publisher caution.',
                    'Strata/non-strata are publisher categories, not universal unit/house definitions. Forty source LGAs have no Strata row; absence stays missing, never zero.',
                    'No raw individual sales, geographic interpolation, invented rates, grand crime totals or cross-state harmonised ranking.',
                    'Current release supplies one original sales quarter and one original rent quarter; published changes are not reconstructed historical medians.',
                    'NSW and VIC crime and rental definitions remain separate.'],
})
print(json.dumps({'sales_lgas': len(sales), 'rent_lgas': len(rent), 'crime_lgas': len(crime),
                  'unmapped_dcj': [identity(n) for n in sorted(dcj_names) if not identity(n)['code']],
                  'crime_excluded': [n for n in sorted(crime_names) if not identity(n)['code']]}))
