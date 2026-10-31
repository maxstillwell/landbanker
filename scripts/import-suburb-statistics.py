"""Reviewed VIC official workbook import. No network or database writes.

Usage: $CODEX_PRIMARY_RUNTIME_PYTHON scripts/import-suburb-statistics.py SOURCE_DIR
Inputs: houses.xlsx, units.xlsx, land.xlsx, crime-lga.xlsx, localities.json.
localities.json is the official Vicmap WFS attribute snapshot (no geometry).
Name-only publisher records join only through this explicit versioned crosswalk
to an exact unique gazetted locality. Ambiguous/neighbourhood records stay unmapped.
"""
import collections
import hashlib
import json
import pathlib
import sys
import openpyxl

source = pathlib.Path(sys.argv[1])
dest = pathlib.Path(__file__).resolve().parents[1] / 'src/lib/layers/data'

def save(name, value):
    (dest / (name + '.json')).write_text(json.dumps(value, ensure_ascii=False, separators=(',', ':'), allow_nan=False) + '\n')

snapshot = json.loads((source / 'localities.json').read_text())
localities = {}
for f in snapshot['features']:
    p = f['properties']
    assert p['locality_name'] not in localities, 'Ambiguous official locality name'
    localities[p['locality_name']] = p
assert len(localities) == snapshot['numberMatched'] == 2973

crosswalk = {}
unmatched = collections.defaultdict(set)
def match(name, product):
    # Whitespace/case only. Never strip directional/locality qualifiers or fuzzy-match.
    name = name.strip().upper()
    if name not in localities:
        unmatched[product].add(name)
        return None
    p = localities[name]
    entry = crosswalk.setdefault(name, {'pfi': p['pfi'], 'vicnames_id': p['vicnames_id'], 'gazetted_name': p['gazetted_locality_name'], 'method': 'exact unique Vicmap bounded locality name', 'products': []})
    if product not in entry['products']:
        entry['products'].append(product)
    return p['pfi']

prices = {}
for kind, filename in [('house', 'houses'), ('unit', 'units'), ('land', 'land')]:
    sheet = openpyxl.load_workbook(source / (filename + '.xlsx'), read_only=True, data_only=True).worksheets[0]
    header = list(next(sheet.iter_rows(min_row=2, max_row=2, values_only=True)))
    assert [header[i] for i in range(1, 22, 2)] == list(range(2015, 2026))
    seen = set()
    for row in sheet.iter_rows(min_row=6, values_only=True):
        if not isinstance(row[0], str) or not row[0].strip():
            continue
        name = row[0].strip()
        assert name not in seen
        seen.add(name)
        code = match(name, 'prices_' + kind)
        if code is None:
            continue
        record = prices.setdefault(code, {})
        for year in range(2015, 2026):
            column = 1 + (year - 2015) * 2
            raw, marker = row[column], str(row[column + 1] or '').strip()
            try:
                value = float(raw)
            except (TypeError, ValueError):
                value = None
            value = value if value is not None and value > 0 else None
            target = record.setdefault(str(year), {})
            target[kind] = value
            target[kind + '_source_marker'] = marker or None
            # No sales volume column exists in these suburb time-series workbooks.
            # Star-marked carried medians are retained in the file, never mapped as observed sales.
    print(filename, 'source localities', len(seen))

crime = {}
seen = set()
workbook = openpyxl.load_workbook(source / 'crime-lga.xlsx', read_only=True, data_only=True)
for row in workbook['Table 03'].iter_rows(min_row=2, max_col=9, values_only=True):
    year, ending, lga, postcode, name, division, subdivision, subgroup, count = row
    if not isinstance(year, int):
        continue
    assert ending == 'June' and 2017 <= year <= 2026
    assert isinstance(count, (int, float)) and count >= 0
    key = (year, lga, postcode, name, subgroup)
    assert key not in seen, 'Duplicate published suburb/subgroup row'
    seen.add(key)
    code = match(name, 'crime')
    if code is None:
        continue
    target = crime.setdefault(code, {}).setdefault(str(year), {'offences': 0, 'source_lgas': [], 'source_postcodes': []})
    target['offences'] += count
    if lga not in target['source_lgas']:
        target['source_lgas'].append(lga)
    if str(postcode) not in target['source_postcodes']:
        target['source_postcodes'].append(str(postcode))
    for metric, include in [('burglary', subdivision.startswith('B30 ')), ('theft', subdivision.startswith('B40 ')), ('motor_theft', subgroup.startswith('B41 ')), ('assault', subdivision.startswith('A20 '))]:
        if include:
            target[metric] = target.get(metric, 0) + count

save('vic-suburb-prices', prices)
save('vic-suburb-crime', crime)
save('vic-suburb-crosswalk', {'version': 'Vicmap bounded locality snapshot 2026-10-31', 'boundary_count': len(localities), 'rows': crosswalk, 'unmatched': {k: sorted(v) for k, v in unmatched.items()}})
base = 'https://www.land.vic.gov.au/__data/assets/excel_doc/'
urls = {
    'houses.xlsx': base + '0033/775617/houses-by-suburb-2015-2025.xlsx',
    'units.xlsx': base + '0034/775618/units-by-suburb-2015-2025.xlsx',
    'land.xlsx': base + '0035/775619/land-by-suburb-2015-2025.xlsx',
    'crime-lga.xlsx': 'https://files.crimestatistics.vic.gov.au/2026-09/Data_Tables_LGA_Recorded_Offences_Year_Ending_June_2026.xlsx',
    'localities.json': 'https://opendata.maps.vic.gov.au/geoserver/wfs?service=WFS&version=2.0.0&request=GetFeature&typeNames=open-data-platform:locality_polygon&outputFormat=application%2Fjson&propertyName=pfi,locality_name,gazetted_locality_name,vicnames_id&count=10000',
}
evidence = {
    'houses.xlsx': 'https://discover.data.vic.gov.au/dataset/victorian-property-sales-report-median-house-by-suburb-time-series',
    'units.xlsx': 'https://discover.data.vic.gov.au/dataset/victorian-property-sales-report-median-unit-by-suburb-time-series',
    'land.xlsx': 'https://discover.data.vic.gov.au/dataset/victorian-property-sales-report-median-vacant-land-by-suburb-time-series',
    'crime-lga.xlsx': 'https://discover.data.vic.gov.au/dataset/data-tables-recorded-offences',
    'localities.json': 'https://discover.data.vic.gov.au/dataset/vicmap-admin-locality-polygon-aligned-to-property',
}
save('suburb-import-manifest', {
    'imported_on': '2026-10-31',
    'sources': [{'file': n, 'url': url, 'sha256': hashlib.sha256((source / n).read_bytes()).hexdigest(), 'licence': 'CC BY 4.0', 'licence_evidence': evidence[n], 'format': 'GeoJSON attributes' if n.endswith('.json') else 'XLSX', 'reference_period': 'Identity snapshot 31 October 2026' if n.endswith('.json') else 'Years ending June 2017–2026' if n.startswith('crime') else 'Calendar years 2015–2025'} for n, url in urls.items()],
    'mapped_price_localities': len(prices), 'mapped_crime_localities': len(crime), 'crime_source_rows': len(seen), 'unmatched_counts': {k: len(v) for k, v in unmatched.items()},
    'methodology_source': 'https://www.land.vic.gov.au/__data/assets/pdf_file/0030/775614/a-guide-to-property-values-2025.pdf',
    'methodology_section': 'Explanatory notes: Symbols and suburb/township classification of sales',
    'commercial_reuse': 'CC BY 4.0 permits commercial reuse and redistribution with attribution, modifications indicated and no endorsement. Original files are not redistributed.',
})
print('mapped', len(prices), len(crime), 'crime rows', len(seen))
